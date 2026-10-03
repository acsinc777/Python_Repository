from datetime import datetime
import csv
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference

# This function creates a chart from the .csv file data
# Argument accepted is the file path and chart type as a string
# No return value
def createChart(path, chartType):
    print("Choose data source to chart. 1.) Inches 2.) Centimeters")
    source = input("Enter a numbered selection.")

    dates = []
    values = []

    try:
        with open(path, 'r') as f:
            for line in f:
                data = line.strip().split(',')
                dates.append(data[0])

                if source == "1":
                    values.append(float(data[1]))
                elif source == "2":
                    values.append(float(data[2]))
                else:
                    print("Invalid data source selected.")
                    return
        workbook = Workbook()
        sheet = workbook.active

        sheet["A1"] = "Date"

        if source == "1":
            sheet["B1"] = "Inches"
        else:
            sheet["B1"] = "Centimeters"
            
        for i in range(len(dates)):
            sheet.cell(row=i + 2, column=1, value=dates[i])
            sheet.cell(row=i + 2, column=2, value=values[i])

        if chartType == "line":
            chart = LineChart()
        elif chartType == "bar":
            chart = BarChart()
        else:
            print("Invalid chart type.")
            return

        data = Reference(sheet, min_col=2, min_row=1, max_row=len(values) + 1)
        categories = Reference(sheet, min_col=1, min_row=2, max_row=len(dates) + 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)

        chart.title = f"ALEDUK5688 {datetime.now().strftime('%m/%d/%Y')}"
        chart.x_axis.title = "Date"

        if source == "1":
            chart.y_axis.title = "Rainfall (Inches)"
        else:
            chart.y_axis.title = "Rainfall (Centimeters)"

        sheet.add_chart(chart, "D2")
        workbook.save("final.xlsx")
        print("Chart successfully created and saved to final.xlsx.")
    except Exception as e:
        print("Error creating chart:", e)

# This function asks the user what kind of graph they want to display
# Argument accepted is file path as a string with no return value
def generateReport(path):
    try:
        print("Choose graph type:")
        print("1.) Line")
        print("2.) Bar")
        choice = input("Enter a numbered selection: ")

        if choice == "1":
            createChart(path, "line")
        elif choice == "2":
            createChart(path, "bar")
        else:
            print("Invalid graph type selected.")

    except Exception as e:
        print("Error generating report:", e)


# This function inserts data into the .csv file
def insertData(path, data):
    try:
        with open(path, 'a') as f:
            f.write(data + "\n")
    except Exception as e:
        print("Error writing to file", e)

# This function reads and displays the contents stored in the .csv file
def viewData(path):
    try:
        with open(path, 'r') as f:
            print("File:", path)
            print(f.read())
    except Exception as e:
        print("Error reading file:", e)

# This function converts inches of rainfall into centimeters
def convertData(inches):
    centi = inches * 2.54
    return centi

# This function takes manual input to collect data
def getInput():
    try:
        counter = int(input("How many entries are you entering?: "))
        for i in range(counter):
            date = (input("Enter a date in the xx/xx/xxxx format: "))
            inches = float(input("Enter the rain amount in inches for the date: "))
            
            # Name of function being called is convertData() with inches being passed into it
            # Expected return value should be centi
            centi = convertData(inches)
            
            print(f'The following was saved at ', str(datetime.now()))
            print(f'{date} | {inches} | {centi}')

            data = f"{date},{inches},{centi}"

            insertData("ZooData.csv", data)
    except Exception as e:
        print("Error entering or saving this data", e)

# This function is the main function
def main():
    print("ALEDUK5688 Spreadsheet Automation Menu")

    menuOptions = ["1.) Input Data", "2.) View Current Data", "3.)Generate Report"]
    count = 0

    #For loop will cycle through the menuOptions list, using count to list them out
    print("Choose one of the numbered options.")
    for count in menuOptions:
        print(count)

    #The next line retrieves the inputted option and stores into the variable called 'Choice'
    choice = int(input("Make your selection: "))
    valid = "Choice selected."

    if choice == 1:
        print(f'Option {choice} was selected at', str(datetime.now()))
        getInput()
    elif choice == 2:
        print(f'Option {choice} was selected at', str(datetime.now()))
        viewData("ZooData.csv")
    elif choice == 3:
        print(f'Option {choice} was selected at', str(datetime.now()))
        generateReport("ZooData.csv")
    else:
        print("Incorrect Option Chosen.")

main()
